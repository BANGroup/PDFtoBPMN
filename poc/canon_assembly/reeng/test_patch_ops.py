"""Тесты patch_ops (синтетический docx, без корпуса и Word): pytest reeng/test_patch_ops.py"""
import os, sys, json, copy
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pytest, docx
from docx.oxml.ns import qn
import patch_ops as PO
import verify_canon

TEXTS = ['Общие положения документа для проверки правок.', 'Первый пункт канона, после него вставляется новый абзац.',
         'Повторяющийся абзац канона для проверки якоря.', 'Повторяющийся абзац канона для проверки якоря.',
         'Удаляемый абзац канона, которого нет в эталоне.', 'Заменяемый абзац канона со старым текстом.', 'Последний абзац канона.']
OK = {'lines': 10, 'missing': 0, 'numbered': 0, 'numbered_bad': 0, 'duplicates': 0, 'frozen_numbers': 0, 'live_numbering': 0, 'coverage': 1.0}


@pytest.fixture
def unit(tmp_path, monkeypatch):
    monkeypatch.setattr(verify_canon, 'verify', lambda *a, **k: dict(OK))   # сама проверка здесь не предмет теста
    od = tmp_path / 'out'; od.mkdir(); sd = tmp_path / 'src'; (sd / 'word').mkdir(parents=True)
    d = docx.Document()
    for t in TEXTS: d.add_paragraph(t)
    d.save(str(od / 'canon.docx')); (sd / 'word' / 'base.docx').write_bytes(b'base')
    return str(od), str(sd)


def texts(od): return [p.text for p in docx.Document(os.path.join(od, 'canon.docx')).paragraphs]
def run(unit, ops, **kw):
    od, sd = unit
    return PO.apply(od, {'ops': ops}, sd, **kw)


def _live_doc(od):
    import pdf_frag as F
    d = docx.Document(os.path.join(od, 'canon.docx')); lists = F.Lists(d); nid = lists.get('dec', 'ru')
    for k, il in ((0, 0), (1, 1), (2, 1)): F.set_numpr(d.paragraphs[k], nid, il)   # 1 / 1.1 / 1.2
    d.save(os.path.join(od, 'canon.docx')); return nid


def test_insert_number_typed_without_live_neighbors(unit):
    od, sd = unit
    r = run(unit, [{'op': 'insert_after', 'anchor': TEXTS[1], 'text': 'Новый пункт из PDF.', 'number': '1.2', 'reason': 'x'}])
    assert len(r['applied']) == 1 and not r['skipped']
    d = docx.Document(os.path.join(od, 'canon.docx')); p = d.paragraphs[2]
    assert p.text == '1.2 Новый пункт из PDF.' and p._p.find('.//' + qn('w:numPr')) is None     # соседи без нумерации: номер текстом
    assert 'PATCH_1' in [b.get(qn('w:name')) for b in p._p.iter(qn('w:bookmarkStart'))]
    assert os.path.exists(os.path.join(od, 'canon_text.txt')) and os.path.exists(os.path.join(od, 'verify.json'))


def test_insert_number_continues_live_list(unit):
    import pynum
    od, sd = unit; nid = _live_doc(od)
    r = run(unit, [{'op': 'insert_after', 'anchor': TEXTS[1], 'text': 'Новый пункт из PDF.', 'number': '1.2'}])
    assert len(r['applied']) == 1 and not r['skipped'], r
    d = docx.Document(os.path.join(od, 'canon.docx'))
    assert d.paragraphs[2].text == 'Новый пункт из PDF.' and d.paragraphs[2]._p.find('.//' + qn('w:numId')).get(qn('w:val')) == str(nid)
    ls = [l for _, l in pynum.list_strings(os.path.join(od, 'canon.docx'))][:4]
    assert ls == ['1', '1.1', '1.2', '1.3']                         # следующий пункт того же списка сдвинулся (допустимо)
    # номер не сходится с продолжением списка: своя цепочка списка (клон abstractNum со start, счёт от 1.9), не набранный текст
    r = run(unit, [{'op': 'insert_after', 'anchor': TEXTS[1], 'text': 'Другой пункт.', 'number': '1.9'}])
    assert len(r['applied']) == 1 and 'цепочка' in r['applied'][0]['info'], r
    assert docx.Document(os.path.join(od, 'canon.docx')).paragraphs[2].text == 'Другой пункт.'
    assert [l for _, l in pynum.list_strings(os.path.join(od, 'canon.docx'))][2] == '1.9'
    # и с PATCH_LIST_START=0 — как раньше, текстом
    os.environ['PATCH_LIST_START'] = '0'
    try:
        r = run(unit, [{'op': 'insert_after', 'anchor': TEXTS[1], 'text': 'Третий пункт.', 'number': '1.7'}])
    finally: del os.environ['PATCH_LIST_START']
    assert docx.Document(os.path.join(od, 'canon.docx')).paragraphs[2].text == '1.7 Третий пункт.'


def test_side_effects_detected():
    import patch_ops as PO
    class E: pass
    a, b, c, new = E(), E(), E(), E()
    before = {a: ('1', '1', 'А'), b: ('1', '2', 'Б'), c: ('2', '5', 'В')}
    ok = {a: ('1', '1', 'А'), new: ('1', '2', 'Н'), b: ('1', '3', 'Б'), c: ('2', '5', 'В')}
    assert PO.side_effects(before, ok, [new], []) == []
    assert 'изменён номер' in PO.side_effects(before, {**ok, c: ('2', '6', 'В')}, [new], [])[0]       # другой список
    assert 'изменён текст' in PO.side_effects(before, {**ok, c: ('2', '5', 'В+')}, [new], [])[0]
    assert 'пропал' in PO.side_effects(before, {k: v for k, v in ok.items() if k is not c}, [new], [])[0]


def test_insert_before_by_prefix_anchor(unit):
    r = run(unit, [{'op': 'insert_before', 'anchor': TEXTS[6][:12] + TEXTS[6][12:], 'text': 'Перед последним.'}])
    assert len(r['applied']) == 1 and texts(unit[0])[-2] == 'Перед последним.'
    r = run(unit, [{'op': 'insert_before', 'anchor': 'Удаляемый абзац канона, которого нет в эт', 'text': 'Перед удаляемым.'}])   # начало >= 30 знаков
    assert len(r['applied']) == 1 and 'Перед удаляемым.' in texts(unit[0])


def test_delete_and_replace(unit):
    od, _ = unit
    r = run(unit, [{'op': 'delete', 'anchor': TEXTS[4]}, {'op': 'replace', 'anchor': TEXTS[5], 'text': 'Новый текст абзаца.'}])
    assert len(r['applied']) == 2 and not r['skipped']
    t = texts(od)
    assert TEXTS[4] not in t and TEXTS[5] not in t and 'Новый текст абзаца.' in t and len(t) == len(TEXTS) - 1
    names = [b.get(qn('w:name')) for b in docx.Document(os.path.join(od, 'canon.docx')).element.iter(qn('w:bookmarkStart'))]
    assert names == ['PATCH_2']


def test_ambiguous_anchor_not_applied_unless_occurrence(unit):
    od, _ = unit
    r = run(unit, [{'op': 'delete', 'anchor': TEXTS[2]}])
    assert not r['applied'] and 'неоднозначен' in r['skipped'][0][1] and texts(od) == TEXTS
    r = run(unit, [{'op': 'delete', 'anchor': TEXTS[2], 'anchor_occurrence': 2}])
    assert len(r['applied']) == 1 and texts(od).count(TEXTS[2]) == 1
    r = run(unit, [{'op': 'delete', 'anchor': 'Такого абзаца нет в каноне вообще'}, {'op': 'insert_after', 'anchor': TEXTS[0]}])
    assert [w for _, w in r['skipped']] == ['якорь не найден', 'нет text']


def test_set_number(unit):
    od, _ = unit
    r = run(unit, [{'op': 'set_number', 'anchor': TEXTS[6], 'number': '3'}])
    assert len(r['applied']) == 1
    p = docx.Document(os.path.join(od, 'canon.docx')).paragraphs[-1]
    assert p.text == '3 ' + TEXTS[6]


def test_part_filter(unit):
    r = run(unit, [{'op': 'delete', 'anchor': TEXTS[4], 'part': 'part_2'}], part='part_1')
    assert not r['applied'] and not r['skipped']
    r = run(unit, [{'op': 'delete', 'anchor': TEXTS[4]}], part='part_1')
    assert 'part' in r['skipped'][0][1]


def test_stale_sha_not_applied(unit, monkeypatch, tmp_path):
    od, sd = unit
    monkeypatch.setattr(PO, 'PATCH_DIR', str(tmp_path / 'patches')); os.makedirs(PO.PATCH_DIR)
    patch = {'doc': 'D', 'files_sha': PO.doc_files_sha(sd), 'ops': [{'op': 'delete', 'anchor': TEXTS[4]}]}
    json.dump(patch, open(PO.patch_path('D'), 'w'))
    st = {'warnings': []}
    PO.run_hook(sd, od, 'D', None, None, dict(OK), st)
    assert st['patches']['applied'] == 1 and not st['patches']['stale'] and TEXTS[4] not in texts(od)
    # документ изменился: sha не совпадает -> не накладывать
    open(os.path.join(sd, 'word', 'base.docx'), 'wb').write(b'changed')
    d = docx.Document(); [d.add_paragraph(t) for t in TEXTS]; d.save(os.path.join(od, 'canon.docx'))
    st = {'warnings': []}
    PO.run_hook(sd, od, 'D', None, None, dict(OK), st)
    assert st['patches']['stale'] and st['patches']['applied'] == 0 and texts(od) == TEXTS
    assert any('правки устарели' in w for w in st['warnings'])


def test_rollback_when_verify_worse(unit, monkeypatch, tmp_path):
    od, sd = unit
    monkeypatch.setattr(PO, 'PATCH_DIR', str(tmp_path / 'patches')); os.makedirs(PO.PATCH_DIR)
    json.dump({'files_sha': PO.doc_files_sha(sd), 'ops': [{'op': 'delete', 'anchor': TEXTS[4]}]}, open(PO.patch_path('D'), 'w'))
    monkeypatch.setattr(verify_canon, 'verify', lambda *a, **k: dict(OK, missing=3))
    st = {'warnings': []}
    v = PO.run_hook(sd, od, 'D', None, None, dict(OK), st)
    assert st['patches'].get('rolled_back') == 1 and st['patches']['applied'] == 0 and texts(od) == TEXTS and v['missing'] == 0


def test_list_like_inherits_numpr(unit):
    import pdf_frag as F
    od, sd = unit
    d = docx.Document(os.path.join(od, 'canon.docx')); lists = F.Lists(d)
    nid = lists.get('dec', 'ru')
    F.set_numpr(d.paragraphs[1], nid, 2); F.set_numpr(d.paragraphs[2], nid, 3)
    d.save(os.path.join(od, 'canon.docx'))
    def numpr(p):
        e = p._p.find('.//' + qn('w:numPr'))
        return None if e is None else (e.find(qn('w:numId')).get(qn('w:val')), e.find(qn('w:ilvl')).get(qn('w:val')))
    r = run(unit, [{'op': 'insert_after', 'anchor': TEXTS[1], 'text': 'Как якорь.', 'list_like': 'anchor'},
                   {'op': 'insert_before', 'anchor': TEXTS[0], 'text': 'Нет списка у якоря.', 'list_like': 'anchor'},
                   {'op': 'insert_before', 'anchor': TEXTS[1], 'text': 'Как следующий.', 'list_like': 'next'},
                   {'op': 'insert_after', 'anchor': TEXTS[6], 'text': 'x', 'list_like': 'anchor', 'number': '1.1'},
                   {'op': 'insert_after', 'anchor': TEXTS[6], 'text': 'x', 'list_like': 'bad'}])
    assert len(r['applied']) == 3 and len(r['skipped']) == 2
    ps = docx.Document(os.path.join(od, 'canon.docx')).paragraphs
    byt = {p.text: p for p in ps}
    assert numpr(byt['Как якорь.']) == (str(nid), '2') and numpr(byt['Как следующий.']) == (str(nid), '2')
    assert numpr(byt['Нет списка у якоря.']) is None


def test_ooxml_rollback(unit, monkeypatch):
    import ooxml_check
    od, sd = unit
    calls = []
    monkeypatch.setattr(ooxml_check, 'check', lambda p: calls.append(p) or ([] if len(calls) == 1 else ['numbering: порядок в lvl']))
    r = run(unit, [{'op': 'delete', 'anchor': TEXTS[4]}])
    assert not r['applied'] and 'откат: ooxml' in r['skipped'][0][1] and texts(od) == TEXTS


def _styled_doc(od):
    """Стиль «H» несёт нумерацию (numPr в стиле); абзацы 1 и 2 — этого стиля; перед абзацем 2 — скрытый пустой нумерованный."""
    import pdf_frag as F
    from docx.oxml import OxmlElement
    d = docx.Document(os.path.join(od, 'canon.docx')); lists = F.Lists(d); nid = lists.get('dec', 'ru')
    st = d.styles.add_style('H', 1)
    ppr = st.element.get_or_add_pPr(); np_ = OxmlElement('w:numPr')
    for tag, val in (('w:ilvl', '0'), ('w:numId', str(nid))):
        e = OxmlElement(tag); e.set(qn('w:val'), val); np_.append(e)
    ppr.append(np_)
    d.paragraphs[1].style = st; d.paragraphs[3].style = st
    hid = copy.deepcopy(d.paragraphs[3]._p); d.paragraphs[3]._p.addprevious(hid)
    from docx.text.paragraph import Paragraph
    h = Paragraph(hid, d.paragraphs[3]._parent); h.text = ''
    rpr = OxmlElement('w:rPr'); rpr.append(OxmlElement('w:vanish')); hid.get_or_add_pPr().append(rpr)
    d.save(os.path.join(od, 'canon.docx'))


def test_set_number_on_style_numbered_paragraph_skipped(unit):
    od, _ = unit; _styled_doc(od)
    r = run(unit, [{'op': 'set_number', 'anchor': TEXTS[1], 'number': '8.1'}])
    assert not r['applied'] and 'стиль абзаца несёт нумерацию' in r['skipped'][0][1]
    assert texts(od)[1] == TEXTS[1]


def test_delete_hidden_before(unit):
    od, _ = unit; _styled_doc(od)
    before = texts(od)
    assert before.count('') == 1
    r = run(unit, [{'op': 'delete_hidden_before', 'anchor': TEXTS[3], 'anchor_occurrence': 2}])
    assert len(r['applied']) == 1, r
    assert texts(od).count('') == 0 and len(texts(od)) == len(before) - 1
    r = run(unit, [{'op': 'delete_hidden_before', 'anchor': TEXTS[3], 'anchor_occurrence': 2}])   # больше скрытых нет
    assert not r['applied'] and 'нет скрытых' in r['skipped'][0][1]


def test_stale_by_content_not_by_name(unit):
    od, sd = unit
    patch = {'files_sha': PO.doc_files_sha(sd)}
    os.rename(os.path.join(sd, 'word', 'base.docx'), os.path.join(sd, 'word', 'base_AB12CD34.docx'))   # переименование — не изменение
    assert PO.stale_reasons(sd, patch) == []
    open(os.path.join(sd, 'word', 'base_AB12CD34.docx'), 'wb').write(b'other')
    assert PO.stale_reasons(sd, patch)


def test_unnumber(unit):
    import pynum
    od, _ = unit; _styled_doc(od)
    ls = lambda: {pynum.para_text(el): l for el, l in pynum.list_strings(os.path.join(od, 'canon.docx')) if pynum.para_text(el)}
    assert ls()[TEXTS[1]] != ''                                         # абзац 1 получил номер от стиля
    r = run(unit, [{'op': 'unnumber', 'anchor': TEXTS[1]}])
    assert len(r['applied']) == 1 and not r['skipped'], r
    assert ls()[TEXTS[1]] == '' and TEXTS[1] in texts(od)


def _table_doc(od):
    d = docx.Document(os.path.join(od, 'canon.docx'))
    t = d.add_table(rows=2, cols=2); t.cell(0, 0).text = 'Старое оглавление раздела'; t.cell(0, 1).text = 'Стр.'
    t.cell(1, 0).text = 'Пункт'; t.cell(1, 1).text = '3'
    d.add_paragraph('Абзац после таблицы.')
    d.save(os.path.join(od, 'canon.docx'))


def test_delete_table(unit):
    od, _ = unit; _table_doc(od)
    r = run(unit, [{'op': 'delete_table', 'anchor': 'Пункт'}, {'op': 'delete_table', 'anchor': 'Старое оглавление раздела'}])
    assert len(r['applied']) == 1 and 'первый абзац таблицы' in r['skipped'][0][1], r
    d = docx.Document(os.path.join(od, 'canon.docx'))
    assert not d.tables and d.paragraphs[-1].text == 'Абзац после таблицы.'
    r = run(unit, [{'op': 'delete_table', 'anchor': TEXTS[0]}])
    assert 'не в ячейке' in r['skipped'][0][1]


def test_delete_hidden_before_through_bookmarks(unit):
    od, _ = unit; _styled_doc(od)
    d = docx.Document(os.path.join(od, 'canon.docx'))
    anchor = d.paragraphs[4]._p; hid = anchor.getprevious()
    from docx.oxml import OxmlElement
    bm = OxmlElement('w:bookmarkEnd'); bm.set(qn('w:id'), '77'); anchor.addprevious(bm)    # закладка между скрытым и якорем
    d.save(os.path.join(od, 'canon.docx'))
    r = run(unit, [{'op': 'delete_hidden_before', 'anchor': TEXTS[3], 'anchor_occurrence': 2}])
    assert len(r['applied']) == 1 and not r['skipped'], r


def test_toc_bookmarks_move_to_new_heading(unit):
    od, _ = unit
    d = docx.Document(os.path.join(od, 'canon.docx')); p = d.paragraphs[4]._p
    from docx.oxml import OxmlElement
    s_ = OxmlElement('w:bookmarkStart'); s_.set(qn('w:id'), '55'); s_.set(qn('w:name'), '_Toc123'); e_ = OxmlElement('w:bookmarkEnd'); e_.set(qn('w:id'), '55')
    p.insert(1, s_); p.append(e_); d.save(os.path.join(od, 'canon.docx'))
    r = run(unit, [{'op': 'delete', 'anchor': TEXTS[4]}, {'op': 'insert_after', 'anchor': TEXTS[1], 'text': 'Новый заголовок.'}])
    assert len(r['applied']) == 2, r
    d = docx.Document(os.path.join(od, 'canon.docx'))
    new = next(p for p in d.paragraphs if p.text == 'Новый заголовок.')._p
    assert '_Toc123' in [b.get(qn('w:name')) for b in new.iter(qn('w:bookmarkStart'))]
    assert len([b for b in d.element.iter(qn('w:bookmarkEnd')) if b.get(qn('w:id')) == '55']) == 1
