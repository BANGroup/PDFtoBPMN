"""Фрагмент .docx из текстового слоя страниц PDF-эталона (страницы, для которых в Word нет источника; решение human 04.10.2026).

Без ИИ: PyMuPDF (строки с геометрией) + python-docx. Колонтитулы, штампы «Изменение №», «Стр. N из M», строки оглавления (точки-заполнители)
и битый текстовый слой отбрасываются теми же фильтрами, что pagediff.usable_lines / verify_canon. Строки склеиваются в абзацы по геометрии;
номера пунктов («7.2.5 Текст», «1) …», «а) …», «– …») становятся живой нумерацией Word (дальше numfix выравнивает её по эталону).
Двуязычные страницы: русские и английские абзацы идут в разных списках (счётчики не мешают друг другу).
Источник помечается закладкой PDF_src_p<N> (N — номер страницы PDF), охватывающей абзацы страницы.
"""
import re, os, copy, random
import fitz
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pagediff import key, garbage

HDR = re.compile(r'Стр\.?\s*/?\s*(page)?\s*\d|Изменение\s*(/\s*Revision)?\s*№')
PAGENUM = re.compile(r'Стр\.?\s*/?\s*(page)?\s*\d')
STAMP = re.compile(r'(Дата введения|Основание\s*(:|_|Приказ)|Изменение\s*(/\s*Revision)?\s*№)', re.I)
TOC_LINE = re.compile(r'[.…]{5,}\s*\d{1,3}\s*$')
TERM = re.compile(r'[.:;!?»”)]\s*$')

DOT = re.compile(r'^(\d{1,3}(?:\.\d{1,3}){1,7})\.?\s+(?=\S)')          # 7.2.5 Текст
LV0 = re.compile(r'^(\d{1,3})(\.)?\s+(?=\S)')                         # 7 ЗАГОЛОВОК / 1. Пункт
PDEC = re.compile(r'^(\d{1,3})\)\s*(?=\S)')                           # 1) пункт
PLET = re.compile(r'^([а-яa-z])\)\s*(?=\S)')                          # а) пункт
DASH = re.compile(r'^[-–—−‒•·▪■□◦]\s+(?=\S)')                        # – пункт
UPPER_START = re.compile(r'[А-ЯЁA-Z«"(\[]')
RU = 'абвгдежзийклмнопрстуфхцчшщъыьэюя'


def usable_idx(texts, head):
    """Индексы строк страницы, остающиеся после фильтров pagediff/reeng_plan.usable_lines (то же правило) + строки оглавления."""
    j = max([i for i, l in enumerate(texts[:12]) if HDR.search(l)] or [-1])
    out = []
    for i in range(j + 1, len(texts)):
        l = texts[i]
        if key(l) in head or garbage(l): continue
        if PAGENUM.search(l) or STAMP.match(l.strip()): continue
        if TOC_LINE.search(l): continue
        out.append(i)
    return out


def page_lines(pg):
    out = []
    for bi, b in enumerate(pg.get_text('dict')['blocks']):
        for l in b.get('lines', []):
            t = ''.join(s['text'] for s in l['spans']).strip()
            if t:
                x0, y0, x1, y1 = l['bbox']
                out.append({'t': t, 'x0': x0, 'y0': y0, 'x1': x1, 'y1': y1, 'blk': bi})
    return out


def merge_visual(ls):
    """Куски одной визуальной строки (разбиение по словам в justified-тексте) внутри одного блока -> одна строка."""
    out = []
    for l in ls:
        p = out[-1] if out else None
        if p and p['blk'] == l['blk'] and abs(p['y0'] - l['y0']) <= 2.5 and abs(p['y1'] - l['y1']) <= 2.5 and l['x0'] >= p['x1'] - 2:
            p['t'] += ' ' + l['t']; p['x1'] = max(p['x1'], l['x1']); p['idxs'] = p.get('idxs', [p.get('idx')]) + [l.get('idx')]; continue
        out.append(dict(l))
    return out


def marker(t, prev_open):
    """(вид, номер-компоненты, текст без маркера) | None. prev_open — есть незавершённый предыдущий абзац."""
    m = DOT.match(t)
    if m:
        comps = [int(x) for x in m.group(1).split('.')]
        rest = t[m.end():]
        if comps[0] == 0 or m.group(1)[0] == '0': return None   # «03.00 UTC» — время, не номер пункта
        if len(comps) >= 3 or not prev_open or UPPER_START.match(rest): return ('dec', comps, rest)
        return None
    m = PDEC.match(t)
    if m: return ('pdec', [int(m.group(1))], t[m.end():])
    m = PLET.match(t)
    if m and m.group(1) in RU: return ('plet', [RU.index(m.group(1)) + 1], t[m.end():])
    if DASH.match(t): return ('dash', [], t[DASH.match(t).end():])
    m = LV0.match(t)
    if m:
        rest = t[m.end():]
        letters = [c for c in rest[:40] if c.isalpha()]
        caps = bool(letters) and sum(c.isupper() for c in letters) / len(letters) >= 0.6
        if int(m.group(1)) >= 1 and (m.group(2) and UPPER_START.match(rest) or caps) and (caps or not prev_open or UPPER_START.match(rest)):
            return ('dec', [int(m.group(1))], rest)
    return None


def script_of(t):
    cyr = len(re.findall(r'[А-Яа-яЁё]', t)); lat = len(re.findall(r'[A-Za-z]', t))
    return 'en' if lat > cyr else 'ru'


def build_paras(ls):
    """Визуальные строки (с геометрией, в порядке PDF) -> абзацы [{'t','kind','comps','script','y0'}]."""
    if not ls: return []
    col_right = {}
    def right_of(l):
        c = [x['x1'] for x in ls if abs(x['x0'] - l['x0']) <= 45 and x['blk'] == l['blk']] or [l['x1']]
        c2 = [x['x1'] for x in ls if abs(x['x0'] - l['x0']) <= 45]
        return max(max(c), sorted(c2)[int(0.9 * (len(c2) - 1))])
    paras = []
    cur = None; prev = None
    for l in ls:
        t = l['t']
        mk = marker(t, cur is not None and not TERM.search(cur['t']))
        new = cur is None or mk is not None
        if not new:
            pl = prev
            h = max(pl['y1'] - pl['y0'], 8)
            same_col = abs(l['x0'] - pl['x0']) <= 45 and l['y0'] >= pl['y0'] - 1
            if not same_col or l['blk'] != pl['blk'] and l['y0'] < pl['y0'] + 2: new = True
            elif l['y0'] - pl['y1'] > 0.6 * h: new = True                      # пустая строка между абзацами
            elif TERM.search(cur['t']) and pl['x1'] < right_of(pl) - 0.12 * max(right_of(pl) - pl['x0'], 1): new = True   # конец абзаца: короткая последняя строка
        if new:
            body = mk[2] if mk else t
            cur = {'t': body, 'kind': mk[0] if mk else None, 'comps': mk[1] if mk else [], 'y0': l['y0'], 'x0': l['x0'], 'blk': l['blk'], 'idxs': list(l.get('idxs', [l.get('idx')]))}
            paras.append(cur)
        else:
            cur['idxs'] += l.get('idxs', [l.get('idx')])
            if cur['t'].endswith(('-', '–')) and t[:1].islower() and len(cur['t']) > 1 and cur['t'][-2].isalpha(): cur['t'] += t
            else: cur['t'] += ' ' + t
        prev = l
    for p in paras: p['script'] = script_of(p['t'])
    return paras


# ---------- нумерация ----------
def _abstract(aid, kind, starts=None):
    A = OxmlElement('w:abstractNum'); A.set(qn('w:abstractNumId'), str(aid))
    e = OxmlElement('w:nsid'); e.set(qn('w:val'), '%08X' % random.randint(0x10000000, 0x7FFFFFFF)); A.append(e)
    mt = OxmlElement('w:multiLevelType'); mt.set(qn('w:val'), 'multilevel' if kind == 'dec' else 'singleLevel'); A.append(mt)
    e = OxmlElement('w:tmpl'); e.set(qn('w:val'), '%08X' % random.randint(0x10000000, 0x7FFFFFFF)); A.append(e)
    for il in range(9 if kind == 'dec' else 1):
        L = OxmlElement('w:lvl'); L.set(qn('w:ilvl'), str(il))
        def sub(tag, val):
            e = OxmlElement(tag); e.set(qn('w:val'), val); L.append(e)
        sub('w:start', str(starts[il]) if starts and il < len(starts) else '1')
        sub('w:numFmt', {'dec': 'decimal', 'pdec': 'decimal', 'plet': 'russianLower', 'dash': 'bullet'}[kind])
        sub('w:suff', 'space')
        sub('w:lvlText', {'dec': '.'.join('%%%d' % (k + 1) for k in range(il + 1)), 'pdec': '%1)', 'plet': '%1)', 'dash': '–'}[kind])
        sub('w:lvlJc', 'left')
        ppr = OxmlElement('w:pPr'); ind = OxmlElement('w:ind'); ind.set(qn('w:left'), '0'); ind.set(qn('w:firstLine'), '0'); ppr.append(ind); L.append(ppr)
        A.append(L)
    return A


class Lists:
    """numId для пары (вид, язык): один список на пару; создаются при первом использовании."""
    def __init__(self, doc):
        self.root = doc.part.numbering_part.element
        ids = [int(x.get(qn('w:abstractNumId'))) for x in self.root.findall(qn('w:abstractNum'))] + [int(x.get(qn('w:numId'))) for x in self.root.findall(qn('w:num'))]
        self.next = max(ids + [0]) + 100
        self.map = {}; self.prev = {}; self.d0 = {}
    def get(self, kind, script, starts=None, force=False):
        k = (kind, script)
        if force or k not in self.map:
            self.next += 1; aid = nid = self.next
            A = _abstract(aid, kind, starts)
            first_num = self.root.find(qn('w:num'))
            if first_num is not None: first_num.addprevious(A)
            else: self.root.append(A)
            N = OxmlElement('w:num'); N.set(qn('w:numId'), str(nid))
            a = OxmlElement('w:abstractNumId'); a.set(qn('w:val'), str(aid)); N.append(a)
            self.root.append(N)
            self.map[k] = nid
        return self.map[k]

    def seq(self, kind, script, comps):
        """restart=True (pdf_canon): номер абзаца должен совпасть с номером PDF. Продолжаем список, если очередной номер — «следующий» после предыдущего
        того же вида/языка (Word посчитает его так же), иначе новый список (abstractNum) со стартами = номер PDF (нумерация приложений, подпунктов «а)», «1)» начинается заново)."""
        k = (kind, script); prev = self.prev.get(k)
        if kind == 'dash': return self.get(kind, script)
        d = len(comps)
        if prev is None: exp = None
        elif kind == 'dec':
            exp = prev[:d - 1] + [prev[d - 1] + 1] if d <= len(prev) else (prev + [1] if d == len(prev) + 1 else None)
        else: exp = [prev[0] + 1]
        self.prev[k] = list(comps)
        if exp == comps and k in self.map and d >= self.d0[k]: return self.map[k]   # d < d0: сброс уровня со «стартовым» значением дал бы чужой номер -> новый список
        self.d0[k] = d
        return self.get(kind, script, starts=comps, force=True)


def set_numpr(p, nid, il):
    ppr = p._p.get_or_add_pPr()
    np_ = OxmlElement('w:numPr')
    e = OxmlElement('w:ilvl'); e.set(qn('w:val'), str(il)); np_.append(e)
    e = OxmlElement('w:numId'); e.set(qn('w:val'), str(nid)); np_.append(e)
    ppr.append(np_)
    ind = OxmlElement('w:ind'); ind.set(qn('w:left'), '0'); ind.set(qn('w:firstLine'), '0'); ppr.append(ind)


def add_bookmark(p_first, p_last, bid, name):
    s = OxmlElement('w:bookmarkStart'); s.set(qn('w:id'), str(bid)); s.set(qn('w:name'), name)
    e = OxmlElement('w:bookmarkEnd'); e.set(qn('w:id'), str(bid))
    ppr = p_first._p.find(qn('w:pPr'))
    if ppr is not None: ppr.addnext(s)
    else: p_first._p.insert(0, s)
    p_last._p.append(e)


def _tables(pg):
    try:
        if not pg.get_drawings(): return []   # таблицы ищутся по линиям: без векторной графики их нет (find_tables ~0.1 с на страницу)
        return [t for t in pg.find_tables().tables if t.row_count >= 2 and t.col_count >= 2]
    except Exception: return []


def make_fragment(pdf, page_nums, dst, head, sect=None, tables=True, restart=False):
    """docx-фрагмент из страниц PDF (1-based номера) -> статистика {'pages','lines','missed','paras','numbered','tables','empty_pages','bookmarks'}."""
    d = fitz.open(pdf)
    doc = docx.Document()
    body = doc.element.body
    for ch in list(body):
        if ch.tag != qn('w:sectPr'): body.remove(ch)
    lists = Lists(doc)
    st = {'pages': [], 'lines': 0, 'missed': [], 'paras': 0, 'numbered': 0, 'tables': 0, 'empty_pages': [], 'bookmarks': []}
    all_text = []; full = ''
    bid = 100; n_done = 0
    for n in page_nums:
        pg = d[n - 1]
        ls0 = page_lines(pg)
        keep = [ls0[i] for i in usable_idx([l['t'] for l in ls0], head)]
        if not keep: st['empty_pages'].append(n); continue
        tabs = _tables(pg) if tables else []
        tb_text, shown = [], set()
        def in_tab(l):
            cx, cy = (l['x0'] + l['x1']) / 2, (l['y0'] + l['y1']) / 2
            for ti, t in enumerate(tabs):
                x0, y0, x1, y1 = t.bbox
                if x0 - 1 <= cx <= x1 + 1 and y0 - 1 <= cy <= y1 + 1: return ti
            return None
        first_par = last_par = None
        seq = []   # элементы страницы по порядку: ('p', абзац) | ('t', индекс таблицы)
        buf = []
        def flush():
            for p in build_paras(merge_visual(buf)): seq.append(('p', p))
            buf.clear()
        for l in keep:
            ti = in_tab(l)
            if ti is None: buf.append(l)
            elif ti not in shown:
                flush(); seq.append(('t', ti)); shown.add(ti)
        flush()
        page_start = len(body)
        for kind, obj in seq:
            if kind == 'p':
                par = doc.add_paragraph(obj['t'])
                if obj['kind']:
                    il = min(len(obj['comps']) - 1, 8) if obj['kind'] == 'dec' else 0
                    set_numpr(par, lists.seq(obj['kind'], obj['script'], obj['comps']) if restart else lists.get(obj['kind'], obj['script']), il); st['numbered'] += 1
                st['paras'] += 1; all_text.append(obj['t'])
                first_par = first_par or par; last_par = par
            else:
                t = tabs[obj]; rows = t.extract()
                tbl = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows)); tbl.style = 'Table Grid'
                for ri, r in enumerate(rows):
                    for ci, c in enumerate(r):
                        txt = re.sub(r'\s*\n\s*', ' ', c or '').strip(); all_text.append(txt)
                        tbl.cell(ri, ci).text = txt
                st['tables'] += 1
        # закладка: от первого до последнего абзаца страницы
        if first_par is not None:
            bid += 1; nm = f'PDF_src_p{n}'
            add_bookmark(first_par, last_par, bid, nm); st['bookmarks'].append(nm)
        st['pages'].append(n); st['lines'] += len(keep)
        full += key(' '.join(all_text[n_done:])); n_done = len(all_text)
        for l in keep:
            k = key(l['t']); k2 = key(re.sub(r'^\s*(\d+(\.\d+)*\.?|[а-яa-z]\)|\d+\)|[-–•·▪])\s*', '', l['t']))
            if len(k) >= 6 and k not in full and k2 not in full: st['missed'].append((n, l['t'][:60]))
    # размер/поля раздела: как у окружающего раздела базы (колонтитулов нет)
    sp = doc.element.body.find(qn('w:sectPr'))
    if sp is not None and sect is not None:
        for ch in list(sp): sp.remove(ch)
        for ch in sect:
            if etree_local(ch) in ('pgSz', 'pgMar', 'cols'): sp.append(copy.deepcopy(ch))
    if st['missed'] and tables and st['tables']:   # таблица потеряла строки -> страницы без таблиц (абзацы)
        return make_fragment(pdf, page_nums, dst, head, sect, tables=False, restart=restart)
    doc.save(dst)
    return st


def etree_local(el): return el.tag.split('}')[1]
