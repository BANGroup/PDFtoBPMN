"""Имитация numfix на фрагменте (без Word): ListString = docx_numbering, выгрузка пишется вручную.
python numsim.py <frag.docx> <эталон.pdf>"""
import sys, os, json, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..')))
import docx
from docx.oxml.ns import qn
import reeng_plan as rp, numfix
import docx_numbering as dn
from pagediff import key
src, pdf = sys.argv[1:3]
tmp = src + '.dump.txt'
d = docx.Document(src); nb = dn.Numbering(d)
rows = []; i = 0
for ki, el in enumerate(d.element.body.iterchildren()):
    ps = [el] if el.tag == qn('w:p') else []
    for p in ps:
        i += 1; t = ''.join(x.text or '' for x in p.iter(qn('w:t')))
        rows.append(f"{i}\tT\t{nb.number(p) or ''}\t{t}\t{i}")
open(tmp, 'w', encoding='utf-8').write('\n'.join(rows))
out = src + '.fixed.docx'
st = numfix.fix(src, tmp, out, pdf)
print({k: v for k, v in st.items() if k not in ('detail', 'unmatched_list')})
for x in st['detail'][:8]: print('  ', x)
print('unmatched', st['unmatched_list'][:6])
d2 = docx.Document(out); nb2 = dn.Numbering(d2)
for p in d2.paragraphs[:40]:
    n = nb2.number(p._p)
    print((n or '-').ljust(10), ''.join(x.text or '' for x in p._p.iter(qn('w:t')))[:90])
