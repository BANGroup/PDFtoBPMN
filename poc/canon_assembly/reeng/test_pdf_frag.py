"""Тесты pdf_frag (синтетический PDF, без корпуса и Word): pytest reeng/test_pdf_frag.py"""
import os, sys, re
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import pytest, fitz, docx
from docx.oxml.ns import qn
import pdf_frag as F
from pagediff import key, canon_pages

HEAD = set()
pytestmark = pytest.mark.skipif(not os.path.exists('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'), reason='нет шрифта DejaVu')


FONT = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'


def make_pdf(path, pages):
    d = fitz.open()
    for lines in pages:
        pg = d.new_page(width=595, height=842)
        pg.insert_font(fontname='dv', fontfile=FONT)
        for (x, y, t) in lines:
            pg.insert_text((x, y), t, fontname='dv', fontsize=10)
    d.save(path)


def paras(path):
    d = docx.Document(path)
    out = []
    for p in d.paragraphs:
        n = p._p.find('.//' + qn('w:numId'))
        out.append((n.get(qn('w:val')) if n is not None else None, p.text))
    return d, out


def test_usable_idx_filters_header_stamp_toc():
    t = ['Название', 'Стр. 5 из 10', 'Изменение № 3', 'Текст страницы номер один', 'Оглавление ........ 12', 'Дата введения изменения 01.01.2020', 'Конец текста страницы']
    assert [t[i] for i in F.usable_idx(t, set())] == ['Текст страницы номер один', 'Конец текста страницы']


def test_marker_kinds():
    assert F.marker('7.2.5 Текст пункта', False)[:2] == ('dec', [7, 2, 5])
    assert F.marker('3.5 кг груза', True) is None                       # десятичное число внутри предложения
    assert F.marker('3.5 Общие положения', True)[0] == 'dec'
    assert F.marker('7 ПОРЯДОК РАБОТЫ', True)[:2] == ('dec', [7])
    assert F.marker('1 января 2020', False) is None
    assert F.marker('а) подпункт', False)[:2] == ('plet', [1])
    assert F.marker('2) подпункт', False)[0] == 'pdec'
    assert F.marker('– пункт', False)[0] == 'dash'
    assert F.marker('0.5 Текст', False) is None


def test_fragment_numbering_paragraphs_bookmark(tmp_path):
    pdf = str(tmp_path / 'a.pdf')
    make_pdf(pdf, [[(50, 40, 'Стр. 1 из 3'), (50, 100, '7.2.5 Первый пункт документа продолжается'), (50, 112, 'на второй строке этого пункта.'),
                    (50, 140, '7.2.6 Второй пункт документа без продолжения.'), (50, 170, 'Просто абзац без номера и разметки текста.')]])
    out = str(tmp_path / 'f.docx')
    st = F.make_fragment(pdf, [1], out, HEAD)
    d, ps = paras(out)
    texts = [t for _, t in ps if t]
    assert texts[0] == 'Первый пункт документа продолжается на второй строке этого пункта.'
    assert ps[0][0] is not None and ps[1][0] == ps[0][0]          # один список, живая нумерация (numId)
    assert not any(re.match(r'\d+\.\d+', t) for t in texts)         # номера не набраны текстом
    assert texts[2].startswith('Просто абзац') and ps[2][0] is None
    assert st['missed'] == [] and st['bookmarks'] == ['PDF_src_p1'] and st['numbered'] == 2
    xml = d.element.xml
    assert 'w:bookmarkStart' in xml and 'PDF_src_p1' in xml
    nums = d.part.numbering_part.element
    lv = nums.findall('.//' + qn('w:lvl'))
    assert any(l.find(qn('w:lvlText')).get(qn('w:val')) == '%1.%2.%3' for l in lv)


def test_two_scripts_two_lists(tmp_path):
    pdf = str(tmp_path / 'b.pdf')
    make_pdf(pdf, [[(50, 100, '9.9.3.1 Русский пункт документа.'), (300, 120, '9.9.3.1 English clause of the document.'),
                    (50, 160, '9.9.3.2 Второй русский пункт.'), (300, 180, '9.9.3.2 Second English clause.')]])
    out = str(tmp_path / 'f.docx'); F.make_fragment(pdf, [1], out, HEAD)
    _, ps = paras(out)
    ids = [n for n, t in ps if t]
    assert len(ids) == 4 and len(set(ids)) == 2 and len({ids[0], ids[2]}) == 1 and len({ids[1], ids[3]}) == 1 and ids[0] != ids[1]


def test_empty_page_skipped(tmp_path):
    pdf = str(tmp_path / 'c.pdf')
    make_pdf(pdf, [[(50, 40, 'Стр. 1 из 3'), (50, 100, 'Содержание ...... 5')]])
    st = F.make_fragment(pdf, [1], str(tmp_path / 'f.docx'), HEAD)
    assert st['empty_pages'] == [1] and st['pages'] == []


# ---------------- pdf_lines: вставка/замена строк класса B в канон ----------------
import pdf_lines as PL


def _canon(path, pars):
    d = docx.Document()
    for t in pars: d.add_paragraph(t)
    d.save(path)


def _run(tmp_path, canon_pars, pdf_lines, bline):
    pdf = str(tmp_path / 'r.pdf'); can = str(tmp_path / 'c.docx')
    make_pdf(pdf, [[(50, 40, 'Стр. 5 из 9')] + pdf_lines] + [[(50, 40, 'Стр. %d из 9' % k), (50, 100, 'Служебная страница номер %d.' % k)] for k in (6, 7, 8)]); _canon(can, canon_pars)
    d = docx.Document(can)
    n, rep = PL.insert_b_lines(d, pdf, [{'label': '5', 'text': bline, 'cls': 'B'}])
    return d, n, rep


A1 = 'Первый уникальный абзац документа для проверки якорей.'
A3 = 'Следующий уникальный абзац после таблицы четыре.'


def test_replace_variant_text(tmp_path):
    d, n, rep = _run(tmp_path, [A1, 'Резервные порции заказываются в зависимости от класса обслуживания согласно таблицы 4.', A3],
                     [(50, 100, A1), (50, 120, 'Резервные порции заказываются в зависимости от'), (50, 132, 'класса обслуживания согласно таблице 4.'), (50, 170, A3)],
                     'класса обслуживания согласно таблице 4.')
    assert n == 1 and rep[0]['status'] == 'вставлено'
    assert d.paragraphs[1].text == 'Резервные порции заказываются в зависимости от класса обслуживания согласно таблице 4.'
    assert 'PDF_src_p' in d.element.xml


def test_insert_missing_line_inline(tmp_path):
    d, n, rep = _run(tmp_path, [A1, 'Резервные порции заказываются в зависимости от', A3],
                     [(50, 100, A1), (50, 120, 'Резервные порции заказываются в зависимости от'), (50, 132, 'класса обслуживания согласно таблице 4.'), (50, 170, A3)],
                     'класса обслуживания согласно таблице 4.')
    assert n == 1
    assert d.paragraphs[1].text == 'Резервные порции заказываются в зависимости от класса обслуживания согласно таблице 4.'
    assert d.paragraphs[0].text == A1 and d.paragraphs[2].text == A3


def test_classify_variants(tmp_path):
    sd = tmp_path / 'sd'; (sd / 'word').mkdir(parents=True)
    _canon(str(sd / 'word' / 'b.docx'), ['отчёта и вкраплени' + 'и\u0306 знаком'])
    out = PL.classify([('1', 'отчѐта'), ('1', 'вкраплений знаком'), ('1', 'совсем другая строка текста')], str(sd))
    # с 05.10 pagediff.key сам приводит NFC и ё/ѐ->е, ѝ->и: варианты написания — это та же строка Word (класс A), а не отдельный «variant»
    assert [o['cls'] for o in out] == ['A', 'A', 'B']


# ---------------- pdf_words: замена 1–3 слов канона словами PDF ----------------
import pdf_words as PW

C1 = 'Программа гарантии качества формируется на следующий год согласно приложению в срок до декабря. Ответственный за формирование – начальник управления качества.'


def test_words_replace_ok():
    r = PW.plan_edits(C1, C1.replace('управления', 'службы'))
    assert r and [(o[3], o[4], o[2]) for o in r[0]] == [('replace', 'управления', 'службы')]


def test_words_reject_interior_delete_and_script_and_head():
    assert PW.plan_edits(C1, C1.replace('согласно ', '')) is None                       # удаление внутри абзаца ненадёжно
    assert PW.plan_edits(C1, C1.replace('управления', 'management')) is None           # другой алфавит
    assert PW.plan_edits(C1, C1.replace('Программа', 'программа')) is None             # регистр не различие (key)
    assert PW.plan_edits(C1, 'а ' + C1[C1.index(' ') + 1:]) is None                     # начало абзаца PDF — продолжение страницы


def test_words_tail_insert_complete_paragraph():
    a = 'Регистрация не может быть завершена до того момента, пока пассажир не подтвердит ознакомление с ограничениями при перевозке'
    r = PW.plan_edits(a, a + ' багажа.')
    assert r and r[0][0][3] == 'insert'
    assert PW.plan_edits(a, a + ' багажа') is None                                      # PDF-абзац без точки — возможно обрыв


# ---------------- py_numfix.plain_head: заголовок с автономером не дублируется ----------------
def _numfix_case(tmp_path, with_head):
    import py_numfix
    pdf = str(tmp_path / 'n.pdf')
    make_pdf(pdf, [[(50, 100, '1 ОБЩИЕ ПОЛОЖЕНИЯ'), (50, 130, '8.6.7 Порядок выдачи допусков'),
                    (50, 160, '8.6.7.1 Текст пункта документа длиннее двадцати четырех знаков.')]])
    d = docx.Document(); lists = F.Lists(d)
    p = d.add_paragraph('ОБЩИЕ ПОЛОЖЕНИЯ'); F.set_numpr(p, lists.get('dec', 'ru'), 0)
    if with_head:
        h = d.add_paragraph('Порядок выдачи допусков'); F.set_numpr(h, lists.get('dec', 'ru'), 2)   # номер автоматический, в тексте его нет
    d.add_paragraph('Текст пункта документа длиннее двадцати четырех знаков.')
    log = py_numfix.run(d, pdf)
    return [x.get('plain_head', 0) for x in log], [q.text for q in d.paragraphs]


def test_plain_head_not_duplicated_for_autonumbered_heading(tmp_path):
    ph, texts = _numfix_case(tmp_path, True)
    assert sum(ph) == 0 and texts.count('Порядок выдачи допусков') == 1


def test_plain_head_still_added_when_missing(tmp_path):
    ph, texts = _numfix_case(tmp_path, False)
    assert sum(ph) == 1 and texts.count('Порядок выдачи допусков') == 1
