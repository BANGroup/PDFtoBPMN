"""Разбор Word в плоский список элементов с привязкой к структуре (TASK-021, шаг 5).

Основа — Docling 2.131 (роль абзаца: заголовок/список/текст, число таблиц и рисунков для сверки);
то, чего Docling не даёт, — свой разбор XML: автонумерация (docx_numbering), ячейки таблиц по gridSpan/vMerge
(объединённые ячейки без дублей), надписи, рисунки/встроенные Visio и Excel с исходным файлом.

parse_docx(path) -> Parsed:
  items: [{text, kind: p|cell|tb|gfx, role, docling, num, tbl, cell, gfx}]  — порядок как в документе;
         text — как load_paras() из apply_ops (номер автонумерации в начале); для gfx — токен-заполнитель
         из знаков частного использования (не влияет на нормализованный текст).
  tables: {tid: {rows, cols, cells: [{r, c, rs, cs}]}}
  gfx:    {gid: {type figure|visio|excel|chart, rel, name, fmt, alt, preview, docx}}
  docling: {tables, pictures, items}   (для перекрёстной проверки)
"""
import os, re, sys, zipfile, difflib, itertools
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lxml import etree
import docx
from docx.oxml.ns import qn
import docx_numbering as N
import assemble as A

W = qn('w:t')
P, TBL, TR, TC = qn('w:p'), qn('w:tbl'), qn('w:tr'), qn('w:tc')
TXBX = qn('w:txbxContent')
FALLBACK = '{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback'
NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
NS_R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS_V = 'urn:schemas-microsoft-com:vml'
NS_O = 'urn:schemas-microsoft-com:office:office'
NS_WP = 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'
NS_C = 'http://schemas.openxmlformats.org/drawingml/2006/chart'
_gid = itertools.count(1)
_tid = itertools.count(1)
TOKENS = {}   # токен-заполнитель -> gid


def token(gid):
    t = '' + chr(0xe100 + gid) + ''
    TOKENS[t] = gid
    return t


def in_fallback(el):
    return any(a.tag == FALLBACK for a in el.iterancestors())


def own_text_fast(p):
    out = []
    for t in p.iter(W):
        skip = False
        for a in t.iterancestors():
            if a is p:
                break
            if a.tag in (TXBX, FALLBACK):
                skip = True
                break
        if not skip:
            out.append(t.text or '')
    return ''.join(out).strip()


def textboxes(p):
    """Абзацы надписей (вне mc:Fallback), верхнего уровня вложенности."""
    res = []
    for tb in p.iter(TXBX):
        if in_fallback(tb):
            continue
        if any(a.tag == TXBX for a in tb.iterancestors() if a is not p and p in a.iterancestors()):
            continue
        for q in tb.iter(P):
            t = own_text_fast(q)
            if t:
                res.append(t)
    return res


class Parsed:
    def __init__(self, path):
        self.path = path
        self.items, self.tables, self.gfx = [], {}, {}
        self.docling = {}
        self.warn = []


def _rels(z):
    root = etree.fromstring(z.read('word/_rels/document.xml.rels'))
    return {r.get('Id'): (r.get('Target'), r.get('Type').rsplit('/', 1)[-1], r.get('TargetMode')) for r in root}


def _target(rel):
    t = rel[0]
    return t.lstrip('/') if t.startswith('/') else 'word/' + t if not t.startswith('word/') else t


def graphics(p, rels, path, ps):
    """Рисунки абзаца (вне mc:Fallback): картинки, OLE (Visio/Excel), диаграммы."""
    res = []
    for el in p.iter():
        if in_fallback(el):
            continue
        tag = el.tag
        if tag == '{%s}blip' % NS_A:
            rid = el.get('{%s}embed' % NS_R)
            rel = rels.get(rid)
            if not rel:
                continue
            alt = ''
            for a in el.iterancestors():
                dp = a.find('.//{%s}docPr' % NS_WP)
                if dp is not None:
                    alt = (dp.get('descr') or dp.get('title') or '').strip()
                    break
            res.append({'type': 'figure', 'rel': _target(rel), 'alt': alt})
        elif tag == qn('w:object'):
            ole = el.find('.//{%s}OLEObject' % NS_O)
            img = el.find('.//{%s}imagedata' % NS_V)
            prev = _target(rels[img.get('{%s}id' % NS_R)]) if img is not None and img.get('{%s}id' % NS_R) in rels else None
            emb = _target(rels[ole.get('{%s}id' % NS_R)]) if ole is not None and ole.get('{%s}id' % NS_R) in rels else None
            prog = (ole.get('ProgID') if ole is not None else '') or ''
            typ = 'visio' if ('visio' in prog.lower() or (emb or '').lower().endswith(('.vsdx', '.vsd'))) else \
                  'excel' if ('excel' in prog.lower() or (emb or '').lower().endswith(('.xlsx', '.xls'))) else 'ole'
            res.append({'type': typ, 'rel': emb, 'preview': prev, 'prog': prog, 'alt': ''})
        elif tag == '{%s}imagedata' % NS_V:
            if any(a.tag == qn('w:object') for a in el.iterancestors()):
                continue
            rel = rels.get(el.get('{%s}id' % NS_R))
            if rel:
                res.append({'type': 'figure', 'rel': _target(rel), 'alt': (el.get('{%s}title' % NS_O) or '')})
        elif tag == '{%s}chart' % NS_C:
            rel = rels.get(el.get('{%s}id' % NS_R))
            res.append({'type': 'chart', 'rel': _target(rel) if rel else None, 'alt': ''})
    for g in res:
        g['name'] = os.path.basename(g['rel'] or '')
        g['fmt'] = os.path.splitext(g['name'])[1].lower().lstrip('.')
        g['docx'] = path
    return res


def parse_docx(path, use_docling=True):
    ps = Parsed(path)
    d = docx.Document(path)
    z = zipfile.ZipFile(path)
    rels = _rels(z)
    nb = N.Numbering(d)

    def fmt_num(num, t):
        if num:
            num = re.sub(r'^(\d+)\.$', r'\1)', num)
            return f'{num} {t}'
        return t

    def add_gfx(p, kind, extra):
        for g in graphics(p, rels, path, ps):
            gid = next(_gid)
            ps.gfx[gid] = g
            g['id'] = gid
            ps.items.append(dict(text=token(gid), kind='gfx', gfx=gid, **extra))

    def do_p(p, extra):
        num = nb.number(p)
        t = own_text_fast(p)
        if t:
            ps.items.append(dict(text=fmt_num(num, t), raw=t, kind='cell' if extra else 'p', num=num, **extra))
        for tb in textboxes(p):
            ps.items.append(dict(text=tb, raw=tb, kind='tb', num=None, **extra))
        add_gfx(p, 'g', extra)

    def do_tbl(tbl):
        tid = next(_tid)
        rows = tbl.findall(TR)
        cells, vopen = [], {}
        ncols = 0
        for r, tr in enumerate(rows):
            c = 0
            for tc in tr.findall(TC):
                pr = tc.find(qn('w:tcPr'))
                cs = 1
                vm = None
                if pr is not None:
                    g = pr.find(qn('w:gridSpan'))
                    cs = int(g.get(qn('w:val'))) if g is not None else 1
                    v = pr.find(qn('w:vMerge'))
                    vm = None if v is None else (v.get(qn('w:val')) or 'continue')
                if vm == 'continue' and c in vopen:
                    vopen[c]['rs'] += 1
                    cell = None
                else:
                    cell = {'r': r, 'c': c, 'rs': 1, 'cs': cs}
                    cells.append(cell)
                    if vm == 'restart':
                        vopen[c] = cell
                    else:
                        vopen.pop(c, None)
                # абзацы ячейки, включая вложенные таблицы (идут в текст этой ячейки)
                if cell is not None:
                    ex = {'tbl': tid, 'cell': len(cells) - 1}
                    for q in tc.iter(P):
                        if any(a.tag in (TXBX, FALLBACK) for a in q.iterancestors() if tc in a.iterancestors()):
                            continue   # абзацы надписей берёт textboxes()
                        do_p(q, ex)
                c += cs
            ncols = max(ncols, c)
        ps.tables[tid] = {'rows': len(rows), 'cols': ncols, 'cells': cells}

    for el in d.element.body.iterchildren():
        if el.tag == P:
            do_p(el, {})
        elif el.tag == TBL:
            do_tbl(el)
    if use_docling:
        try:
            docling_roles(ps)
        except Exception as ex:   # не прятать: пишем в warn и в GAP записи
            ps.warn.append(f'Docling: {type(ex).__name__}: {ex}')
    return ps


_conv = None


def docling_roles(ps):
    global _conv
    if _conv is None:
        from docling.document_converter import DocumentConverter
        _conv = DocumentConverter()
    doc = _conv.convert(ps.path).document
    dl = []   # (label, level, norm text) для абзацев вне таблиц
    ntab = npic = 0
    for it, _ in doc.iterate_items(with_groups=False):
        lab = str(it.label.value)
        if lab == 'table':
            ntab += 1; continue
        if lab == 'picture':
            npic += 1; continue
        par = it.parent.resolve(doc) if getattr(it, 'parent', None) else None
        if par is not None and str(getattr(par, 'label', '')).endswith('table'):
            continue
        for line in [x for x in (getattr(it, 'text', '') or '').split('\n')] or ['']:
            n = A.norm(line)
            if n:
                dl.append((lab, getattr(it, 'level', None), n))
    ps.docling = {'tables': ntab, 'pictures': npic, 'items': len(dl)}
    mine = [(i, A.norm(it['raw'])) for i, it in enumerate(ps.items) if it['kind'] in ('p', 'tb') and A.norm(it.get('raw', ''))]
    sm = difflib.SequenceMatcher(None, [n for _, n in mine], [x[2] for x in dl], autojunk=False)
    hit = 0
    for a, b, size in sm.get_matching_blocks():
        for k in range(size):
            it = ps.items[mine[a + k][0]]
            it['docling'] = (dl[b + k][0], dl[b + k][1])
            hit += 1
    ps.docling['aligned'] = hit
    ps.docling['mine_paragraphs'] = len(mine)


def texts(ps):
    return [it['text'] for it in ps.items]


# ---------------- Visio / Excel / картинки ----------------
def read_visio(data):
    """vsdx (zip) -> {pages: [{name, shapes: [{id, name, text}], links: [{from, to, label}]}]}"""
    import io
    z = zipfile.ZipFile(io.BytesIO(data))
    ns = {'v': 'http://schemas.microsoft.com/office/visio/2012/main'}
    V = '{%s}' % ns['v']
    names = {}
    if 'visio/pages/pages.xml' in z.namelist():
        pr = etree.fromstring(z.read('visio/pages/pages.xml'))
        rl = {}
        if 'visio/pages/_rels/pages.xml.rels' in z.namelist():
            for r in etree.fromstring(z.read('visio/pages/_rels/pages.xml.rels')):
                rl[r.get('Id')] = r.get('Target')
        for pg in pr.iter(V + 'Page'):
            rel = pg.find(V + 'Rel')
            tgt = rl.get(rel.get('{%s}id' % NS_R)) if rel is not None else None
            if tgt:
                names['visio/pages/' + tgt] = pg.get('NameU') or pg.get('Name')
    pages = []
    for pn in sorted([n for n in z.namelist() if re.fullmatch(r'visio/pages/page\d+\.xml', n)],
                     key=lambda s: int(re.search(r'(\d+)', s.rsplit('/', 1)[1]).group(1))):
        root = etree.fromstring(z.read(pn))
        shapes, byid = [], {}
        for s in root.iter(V + 'Shape'):
            tx = s.find(V + 'Text')
            t = re.sub(r'\s+', ' ', ''.join(tx.itertext())).strip() if tx is not None else ''
            byid[s.get('ID')] = {'id': s.get('ID'), 'name': s.get('NameU') or s.get('Name') or '', 'text': t}
        ends = {}
        for c in root.iter(V + 'Connect'):
            ends.setdefault(c.get('FromSheet'), {})[c.get('FromCell')] = c.get('ToSheet')
        links = []
        for cid, e in ends.items():
            b, en = e.get('BeginX'), e.get('EndX')
            if b is None and en is None:
                continue
            links.append({'from': b, 'to': en, 'connector': cid, 'label': byid.get(cid, {}).get('text', '')})
        connectors = set(ends)
        shapes = [s for i, s in byid.items() if s['text'] and i not in connectors or (i in connectors and s['text'])]
        # связи «от → к» с текстами фигур
        for l in links:
            l['from_text'] = byid.get(l['from'], {}).get('text', '') if l['from'] else ''
            l['to_text'] = byid.get(l['to'], {}).get('text', '') if l['to'] else ''
        pages.append({'name': names.get(pn) or pn, 'shapes': shapes, 'links': links,
                      'shapes_total': len(byid)})
    return {'pages': pages}


def read_excel(data):
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    out = []
    for ws in wb.worksheets:
        cells = []
        merged = [(m.min_row, m.min_col, m.max_row - m.min_row + 1, m.max_col - m.min_col + 1) for m in ws.merged_cells.ranges]
        span = {(r, c): (rs, cs) for r, c, rs, cs in merged}
        inside = set()
        for r, c, rs, cs in merged:
            for i in range(r, r + rs):
                for j in range(c, c + cs):
                    if (i, j) != (r, c):
                        inside.add((i, j))
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None or (cell.row, cell.column) in inside:
                    continue
                rs, cs = span.get((cell.row, cell.column), (1, 1))
                cells.append({'r': cell.row - 1, 'c': cell.column - 1, 'rs': rs, 'cs': cs, 'text': str(cell.value).strip()})
        out.append({'name': ws.title, 'cells': cells})
    return out


def emf_text(data):
    """Текстовые фрагменты EMF (EMR_EXTTEXTOUTW и EMR_SMALLTEXTOUT нет) в порядке записей: [строки]."""
    import struct
    out, off = [], 0
    while off + 8 <= len(data):
        t, sz = struct.unpack_from('<II', data, off)
        if sz < 8 or off + sz > len(data): break
        if t == 84 and sz >= 76:     # EMR_EXTTEXTOUTW
            base = off + 8 + 16 + 12
            x, y, nch, offs = struct.unpack_from('<iiII', data, base)
            s = data[off + offs:off + offs + 2 * nch].decode('utf-16le', errors='replace').strip()
            if s: out.append(s)
        off += sz
    return out


def wmf_text(data):
    """Текстовые фрагменты WMF (META_TEXTOUT, META_EXTTEXTOUT), кодировка cp1251."""
    import struct
    off = 22 if data[:4] == b'\xd7\xcd\xc6\x9a' else 0
    off += 18
    out = []
    while off + 6 <= len(data):
        size, fn = struct.unpack_from('<IH', data, off)
        if size < 3 or off + size * 2 > len(data): break
        p = off + 6
        try:
            if fn == 0x0521:      # TEXTOUT
                n = struct.unpack_from('<H', data, p)[0]
                s = data[p + 2:p + 2 + n]
            elif fn == 0x0A32:    # EXTTEXTOUT
                y, x, n, opt = struct.unpack_from('<hhHH', data, p)
                q = p + 8 + (8 if opt & 6 else 0)
                s = data[q:q + n]
            else:
                s = b''
            s = s.decode('cp1251', errors='replace').strip()
            if s: out.append(s)
        except struct.error:
            pass
        off += size * 2
    return out


def image_info(data, fmt):
    """Формат и читаемость картинки: Pillow открывает png/jpg/gif/bmp/tif; wmf/emf — нет."""
    import io
    from PIL import Image
    info = {'format': fmt, 'readable': False}
    try:
        im = Image.open(io.BytesIO(data))
        im.load()
        info.update(readable=True, width=im.width, height=im.height)
    except Exception as ex:
        info['note'] = f'Pillow не открывает ({type(ex).__name__})'
    return info


if __name__ == '__main__':
    ps = parse_docx(sys.argv[1])
    import collections
    print('элементов', len(ps.items), collections.Counter(i['kind'] for i in ps.items))
    print('таблиц', len(ps.tables), 'gfx', collections.Counter(g['type'] for g in ps.gfx.values()), ps.docling, ps.warn)
